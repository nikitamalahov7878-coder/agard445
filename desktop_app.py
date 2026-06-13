from __future__ import annotations
import io
import queue
import threading
import traceback
from dataclasses import dataclass
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText
from openpyxl import load_workbook

from excel_compare import (
    ComparisonConfig,
    InputFormatError,
    compare_employee_presence,
    compare_report_and_many_timesheets,
    compare_report_and_timesheet,
)
from output_naming import filename_match_score, make_sverka_filename

MODE_SINGLE = 'single'
MODE_MULTI = 'multi'
MODE_BATCH = 'batch_pairs'
MODE_EMPLOYEES = 'employees'

PERIOD_FULL = 'full_month'
PERIOD_1_15 = 'first_half'
PERIOD_15_31 = 'second_half'


@dataclass
class CompareRequest:
    mode: str
    period: str
    primary_path: Path
    secondary_path: Path | None
    timesheet_paths: tuple[Path, ...]
    batch_paths: tuple[Path, ...]
    output_path: Path


def _period_bounds(period_mode: str) -> tuple[int | None, int | None]:
    if period_mode == PERIOD_1_15:
        return (1, 15)
    if period_mode == PERIOD_15_31:
        return (15, 31)
    return (None, None)


def _build_config(period_mode: str) -> ComparisonConfig:
    period_start, period_end = _period_bounds(period_mode)
    return ComparisonConfig(
        threshold_minutes=15,
        lunch_break_minutes=60,
        period_start_day=period_start,
        period_end_day=period_end,
    )


def _normalize_output_path(path_text: str) -> Path:
    path = Path(path_text.strip())
    if path.suffix.lower() != '.xlsx':
        path = path.with_suffix('.xlsx')
    return path


def _compare_two_excel_files(
    file_a: Path, file_b: Path, config: ComparisonConfig
) -> tuple[bytes, str]:
    try:
        result = compare_report_and_timesheet(
            file_a.read_bytes(), file_b.read_bytes(), config=config
        )
        return result, make_sverka_filename(file_a.name, file_b.name)
    except InputFormatError:
        result = compare_report_and_timesheet(
            file_b.read_bytes(), file_a.read_bytes(), config=config
        )
        return result, make_sverka_filename(file_a.name, file_b.name)


def _sheet_data_rows(sheet) -> int:
    data_rows = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(cell is not None for cell in row):
            data_rows += 1
    return data_rows


def _ensure_result_note(result_bytes: bytes, mode: str) -> bytes:
    wb = load_workbook(io.BytesIO(result_bytes))
    total_data_rows = sum(_sheet_data_rows(ws) for ws in wb.worksheets)
    if total_data_rows > 0:
        return result_bytes
    title = 'Результат'
    if title in wb.sheetnames:
        ws = wb[title]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(title, 0)
    if mode == MODE_EMPLOYEES:
        note = 'Различий по сотрудникам не найдено.'
    else:
        note = 'Расхождений не найдено.'
    ws.append(['Статус'])
    ws.append([note])
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _result_data_row_count(result_bytes: bytes) -> int:
    wb = load_workbook(io.BytesIO(result_bytes))
    return sum(_sheet_data_rows(ws) for ws in wb.worksheets)


class SverkaApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title('Сверка Excel')
        self.geometry('980x760')
        self.minsize(860, 680)

        self.mode_var = tk.StringVar(value=MODE_SINGLE)
        self.period_var = tk.StringVar(value=PERIOD_FULL)
        self.primary_path_var = tk.StringVar()
        self.secondary_path_var = tk.StringVar()
        self.output_path_var = tk.StringVar()
        self.status_var = tk.StringVar(value='Готово к работе.')

        self.timesheet_paths: list[Path] = []
        self.batch_paths: list[Path] = []
        self._busy = False
        self._worker_queue: queue.Queue = queue.Queue()

        self._build_ui()
        self._apply_mode()
        self.after(150, self._poll_worker_queue)

    def _build_ui(self) -> None:
        self.columnconfigure(0, weight=1)
        self.rowconfigure(5, weight=1)

        # Header
        header = ttk.Frame(self, padding=(16, 16, 16, 8))
        header.grid(row=0, column=0, sticky='nsew')
        header.columnconfigure(0, weight=1)

        ttk.Label(header, text='Локальная сверка Excel', font=('Segoe UI', 16, 'bold')).grid(
            row=0, column=0, sticky='w'
        )
        self.mode_hint_label = ttk.Label(header, wraplength=900, justify='left')
        self.mode_hint_label.grid(row=1, column=0, sticky='w', pady=(8, 0))

        # Mode frame
        mode_frame = ttk.LabelFrame(self, text='Режим', padding=12)
        mode_frame.grid(row=1, column=0, sticky='ew', padx=16, pady=8)
        for index in range(4):
            mode_frame.columnconfigure(index, weight=1)

        ttk.Radiobutton(
            mode_frame,
            text='1 отчёт + 1 табель',
            variable=self.mode_var,
            value=MODE_SINGLE,
            command=self._on_mode_changed,
        ).grid(row=0, column=0, sticky='w')
        ttk.Radiobutton(
            mode_frame,
            text='1 отчёт + много табелей',
            variable=self.mode_var,
            value=MODE_MULTI,
            command=self._on_mode_changed,
        ).grid(row=0, column=1, sticky='w')
        ttk.Radiobutton(
            mode_frame,
            text='Много пар',
            variable=self.mode_var,
            value=MODE_BATCH,
            command=self._on_mode_changed,
        ).grid(row=0, column=2, sticky='w')
        ttk.Radiobutton(
            mode_frame,
            text='Сверка сотрудников',
            variable=self.mode_var,
            value=MODE_EMPLOYEES,
            command=self._on_mode_changed,
        ).grid(row=0, column=3, sticky='w')

        # Period frame
        period_frame = ttk.LabelFrame(self, text='Период', padding=12)
        period_frame.grid(row=2, column=0, sticky='ew', padx=16, pady=8)

        rb_full = ttk.Radiobutton(
            period_frame,
            text='Весь месяц',
            variable=self.period_var,
            value=PERIOD_FULL,
            command=self._refresh_output_path,
        )
        rb_1_15 = ttk.Radiobutton(
            period_frame,
            text='1-15',
            variable=self.period_var,
            value=PERIOD_1_15,
            command=self._refresh_output_path,
        )
        rb_15_31 = ttk.Radiobutton(
            period_frame,
            text='15-31',
            variable=self.period_var,
            value=PERIOD_15_31,
            command=self._refresh_output_path,
        )
        self.period_buttons = [rb_full, rb_1_15, rb_15_31]
        for idx, button in enumerate(self.period_buttons):
            button.grid(row=0, column=idx, sticky='w', padx=(0, 18))

        # Files frame
        files_frame = ttk.LabelFrame(self, text='Файлы', padding=12)
        files_frame.grid(row=3, column=0, sticky='nsew', padx=16, pady=8)
        files_frame.columnconfigure(1, weight=1)

        self.primary_label = ttk.Label(files_frame, text='Отчёт')
        self.primary_label.grid(row=0, column=0, sticky='w', pady=4)
        self.primary_entry = ttk.Entry(files_frame, textvariable=self.primary_path_var)
        self.primary_entry.grid(row=0, column=1, sticky='ew', padx=8, pady=4)
        self.primary_button = ttk.Button(
            files_frame, text='Выбрать...', command=self._choose_primary
        )
        self.primary_button.grid(row=0, column=2, sticky='ew', pady=4)

        self.secondary_label = ttk.Label(files_frame, text='Табель')
        self.secondary_label.grid(row=1, column=0, sticky='w', pady=4)
        self.secondary_entry = ttk.Entry(files_frame, textvariable=self.secondary_path_var)
        self.secondary_entry.grid(row=1, column=1, sticky='ew', padx=8, pady=4)
        self.secondary_button = ttk.Button(
            files_frame, text='Выбрать...', command=self._choose_secondary
        )
        self.secondary_button.grid(row=1, column=2, sticky='ew', pady=4)

        # Multi-file frame (inside files_frame)
        self.multi_frame = ttk.Frame(files_frame)
        self.multi_frame.grid(row=1, column=0, columnspan=3, sticky='nsew', pady=4)
        self.multi_frame.columnconfigure(0, weight=1)
        self.multi_frame.rowconfigure(1, weight=1)

        self.multi_list_label = ttk.Label(self.multi_frame, text='Табели')
        self.multi_list_label.grid(row=0, column=0, sticky='w', pady=(0, 6))

        multi_buttons = ttk.Frame(self.multi_frame)
        multi_buttons.grid(row=0, column=1, sticky='e', pady=(0, 6))
        ttk.Button(multi_buttons, text='Добавить...', command=self._add_timesheets).grid(
            row=0, column=0, padx=(0, 6)
        )
        ttk.Button(
            multi_buttons, text='Удалить', command=self._remove_selected_timesheet
        ).grid(row=0, column=1, padx=(0, 6))
        ttk.Button(multi_buttons, text='Очистить', command=self._clear_timesheets).grid(
            row=0, column=2
        )

        self.timesheet_listbox = tk.Listbox(self.multi_frame, height=8)
        self.timesheet_listbox.grid(row=1, column=0, columnspan=2, sticky='nsew')
        timesheet_scrollbar = ttk.Scrollbar(
            self.multi_frame, orient='vertical', command=self.timesheet_listbox.yview
        )
        timesheet_scrollbar.grid(row=1, column=2, sticky='ns')
        self.timesheet_listbox.configure(yscrollcommand=timesheet_scrollbar.set)

        # Output frame
        output_frame = ttk.LabelFrame(self, text='Результат', padding=12)
        output_frame.grid(row=4, column=0, sticky='ew', padx=16, pady=8)
        output_frame.columnconfigure(1, weight=1)

        ttk.Label(output_frame, text='Файл результата').grid(row=0, column=0, sticky='w')
        ttk.Entry(output_frame, textvariable=self.output_path_var).grid(
            row=0, column=1, sticky='ew', padx=8
        )
        ttk.Button(output_frame, text='Сохранить как...', command=self._choose_output).grid(
            row=0, column=2, padx=(0, 6)
        )
        ttk.Button(
            output_frame, text='Подставить имя', command=self._refresh_output_path
        ).grid(row=0, column=3)

        # Log frame
        log_frame = ttk.LabelFrame(self, text='Статус и журнал', padding=12)
        log_frame.grid(row=5, column=0, sticky='nsew', padx=16, pady=8)
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(1, weight=1)

        ttk.Label(log_frame, textvariable=self.status_var).grid(
            row=0, column=0, sticky='w', pady=(0, 8)
        )
        self.log_text = ScrolledText(log_frame, height=14, wrap='word', state='disabled')
        self.log_text.grid(row=1, column=0, sticky='nsew')

        # Actions
        actions = ttk.Frame(self, padding=(16, 0, 16, 16))
        actions.grid(row=6, column=0, sticky='ew')
        ttk.Button(
            actions, text='Сформировать сверку', command=self._start_compare
        ).pack(side='left')
        ttk.Button(actions, text='Сбросить', command=self._reset_form).pack(
            side='left', padx=(8, 0)
        )

    def _log(self, message: str) -> None:
        self.log_text.configure(state='normal')
        self.log_text.insert('end', message.rstrip() + '\n')
        self.log_text.see('end')
        self.log_text.configure(state='disabled')

    def _on_mode_changed(self) -> None:
        self._apply_mode()
        self._refresh_output_path()

    def _apply_mode(self) -> None:
        mode = self.mode_var.get()
        self._refresh_dynamic_list()

        if mode == MODE_MULTI:
            self.primary_label.configure(text='Отчёт')
            self.multi_list_label.configure(text='Табели')
            self.mode_hint_label.configure(
                text='Режим: один отчёт и несколько табелей. Выберите один файл отчёта и добавьте все табели, которые нужно объединить в одну общую сверку.'
            )
            self.primary_label.grid()
            self.primary_entry.grid()
            self.primary_button.grid()
            self.secondary_label.grid_remove()
            self.secondary_entry.grid_remove()
            self.secondary_button.grid_remove()
            self.multi_frame.grid()
            self._set_period_controls_enabled(True)

        elif mode == MODE_BATCH:
            self.mode_hint_label.configure(
                text='Режим: много пар. Добавьте сразу все файлы отчётов и табелей. Программа сама найдёт подходящие пары и сохранит отдельную сверку по каждой паре в выбранную папку.'
            )
            self.multi_list_label.configure(text='Файлы для пакетной сверки')
            self.primary_label.grid_remove()
            self.primary_entry.grid_remove()
            self.primary_button.grid_remove()
            self.secondary_label.grid_remove()
            self.secondary_entry.grid_remove()
            self.secondary_button.grid_remove()
            self.multi_frame.grid(row=1, column=0, columnspan=3, sticky='nsew', pady=4)
            self._set_period_controls_enabled(True)

        elif mode == MODE_EMPLOYEES:
            self.primary_label.configure(text='Список сотрудников')
            self.secondary_label.configure(text='Табель')
            self.mode_hint_label.configure(
                text='Режим: сверка сотрудников. На вход подаются список сотрудников и табель. На выходе файл с сотрудниками только в табеле, только в списке и общим количеством.'
            )
            self.secondary_label.grid()
            self.secondary_entry.grid()
            self.secondary_button.grid()
            self.multi_frame.grid_remove()
            self.primary_label.grid()
            self.primary_entry.grid()
            self.primary_button.grid()
            self._set_period_controls_enabled(False)

        else:  # MODE_SINGLE
            self.primary_label.configure(text='Отчёт')
            self.secondary_label.configure(text='Табель')
            self.mode_hint_label.configure(
                text='Режим: один отчёт и один табель. Рабочее время считается по отчёту, расхождения учитываются от 15 минут. Можно ограничить сверку половиной месяца.'
            )
            self.primary_label.grid()
            self.primary_entry.grid()
            self.primary_button.grid()
            self.secondary_label.grid()
            self.secondary_entry.grid()
            self.secondary_button.grid()
            self.multi_frame.grid_remove()
            self._set_period_controls_enabled(True)

    def _set_period_controls_enabled(self, enabled: bool) -> None:
        state = 'normal' if enabled else 'disabled'
        for button in self.period_buttons:
            button.configure(state=state)

    def _choose_primary(self) -> None:
        path = filedialog.askopenfilename(
            title='Выберите файл',
            filetypes=[('Excel files', '*.xlsx'), ('All files', '*.*')],
        )
        if not path:
            return
        self.primary_path_var.set(path)
        self._refresh_output_path()

    def _choose_secondary(self) -> None:
        path = filedialog.askopenfilename(
            title='Выберите файл',
            filetypes=[('Excel files', '*.xlsx'), ('All files', '*.*')],
        )
        if not path:
            return
        self.secondary_path_var.set(path)
        self._refresh_output_path()

    def _add_timesheets(self) -> None:
        paths = filedialog.askopenfilenames(
            title='Выберите файлы',
            filetypes=[('Excel files', '*.xlsx'), ('All files', '*.*')],
        )
        if not paths:
            return
        target_list = self._active_path_list()
        existing = {str(p).lower() for p in target_list}
        for item in paths:
            if item.lower() not in existing:
                target_list.append(Path(item))
        self._refresh_dynamic_list()
        self._refresh_output_path()

    def _remove_selected_timesheet(self) -> None:
        selected = list(self.timesheet_listbox.curselection())
        if not selected:
            return
        target_list = self._active_path_list()
        for index in reversed(selected):
            del target_list[index]
        self._refresh_dynamic_list()
        self._refresh_output_path()

    def _clear_timesheets(self) -> None:
        self._active_path_list().clear()
        self._refresh_dynamic_list()
        self._refresh_output_path()

    def _active_path_list(self) -> list[Path]:
        if self.mode_var.get() == MODE_BATCH:
            return self.batch_paths
        return self.timesheet_paths

    def _refresh_dynamic_list(self) -> None:
        self.timesheet_listbox.delete(0, 'end')
        for path in self._active_path_list():
            self.timesheet_listbox.insert('end', str(path))

    def _choose_output(self) -> None:
        initial_file = self._suggest_output_path()
        if self.mode_var.get() == MODE_BATCH:
            path = filedialog.askdirectory(
                title='Куда сохранить пакет сверок',
                initialdir=str(initial_file) if initial_file else '',
                mustexist=False,
            )
        else:
            path = filedialog.asksaveasfilename(
                title='Куда сохранить сверку',
                defaultextension='.xlsx',
                filetypes=[('Excel files', '*.xlsx')],
                initialfile=initial_file.name if initial_file else '',
                initialdir=str(initial_file.parent) if initial_file else '',
            )
        if not path:
            return
        self.output_path_var.set(path)

    def _suggest_output_path(self) -> Path | None:
        mode = self.mode_var.get()
        primary_text = self.primary_path_var.get().strip()
        secondary_text = self.secondary_path_var.get().strip()

        if mode == MODE_BATCH:
            if not self.batch_paths:
                return None
            base_dir = self.batch_paths[0].resolve().parent
            return base_dir / 'Сверки_пакет'

        if mode == MODE_MULTI:
            if not primary_text or not self.timesheet_paths:
                return None
            base_dir = Path(primary_text).resolve().parent
            filename = make_sverka_filename(
                Path(primary_text).name, 'общая', prefix='Сверка_общая'
            )
            return base_dir / filename

        if mode == MODE_EMPLOYEES:
            if not primary_text or not secondary_text:
                return None
            base_dir = Path(primary_text).resolve().parent
            filename = make_sverka_filename(
                Path(primary_text).name,
                Path(secondary_text).name,
                prefix='Сверка_сотрудники',
            )
            return base_dir / filename

        # MODE_SINGLE
        if not primary_text or not secondary_text:
            return None
        base_dir = Path(primary_text).resolve().parent
        filename = make_sverka_filename(Path(primary_text).name, Path(secondary_text).name)
        return base_dir / filename

    def _refresh_output_path(self) -> None:
        suggested = self._suggest_output_path()
        if suggested is not None:
            self.output_path_var.set(str(suggested))

    def _reset_form(self) -> None:
        if self._busy:
            messagebox.showwarning('Сверка выполняется', 'Дождитесь завершения текущей обработки.')
            return
        self.primary_path_var.set('')
        self.secondary_path_var.set('')
        self.output_path_var.set('')
        self.timesheet_paths.clear()
        self.batch_paths.clear()
        self._refresh_dynamic_list()
        self.period_var.set(PERIOD_FULL)
        self.mode_var.set(MODE_SINGLE)
        self._apply_mode()
        self.status_var.set('Готово к работе.')
        self._log('Форма очищена.')

    def _build_request(self) -> CompareRequest:
        mode = self.mode_var.get()
        primary_text = self.primary_path_var.get().strip()
        secondary_text = self.secondary_path_var.get().strip()
        output_text = self.output_path_var.get().strip()

        if not primary_text and mode != MODE_BATCH:
            raise ValueError('Не выбран основной файл.')

        primary_path = Path(primary_text) if primary_text else Path('.')

        if mode != MODE_BATCH and not primary_path.is_file():
            raise ValueError(f'Файл не найден: {primary_path}')

        if mode == MODE_BATCH:
            if len(self.batch_paths) < 2:
                raise ValueError('Для режима много пар нужно добавить минимум 2 файла.')
            secondary_path = None
            timesheet_paths: tuple[Path, ...] = ()
            batch_paths = tuple(self.batch_paths)

        elif mode == MODE_MULTI:
            if not self.timesheet_paths:
                raise ValueError('Не добавлен ни один табель.')
            secondary_path = None
            timesheet_paths = tuple(self.timesheet_paths)
            batch_paths = ()

        else:
            if not secondary_text:
                raise ValueError('Не выбран второй файл.')
            secondary_path = Path(secondary_text)
            if not secondary_path.is_file():
                raise ValueError(f'Файл не найден: {secondary_path}')
            timesheet_paths = ()
            batch_paths = ()

        if not output_text:
            suggested = self._suggest_output_path()
            if suggested is None:
                raise ValueError('Не удалось определить имя выходного файла.')
            output_path = suggested
            self.output_path_var.set(str(suggested))
        else:
            if mode == MODE_BATCH:
                output_path = Path(output_text)
            else:
                output_path = _normalize_output_path(output_text)

        return CompareRequest(
            mode=mode,
            period=self.period_var.get(),
            primary_path=primary_path,
            secondary_path=secondary_path,
            timesheet_paths=timesheet_paths,
            batch_paths=batch_paths,
            output_path=output_path,
        )

    def _start_compare(self) -> None:
        if self._busy:
            return
        try:
            request = self._build_request()
        except Exception as exc:
            messagebox.showerror('Ошибка', str(exc))
            return
        self._busy = True
        self.status_var.set('Идёт обработка файлов...')
        self._log(f'Запуск сверки: режим = {self._mode_name(request.mode)}')
        self._log(f'Сохранение результата: {request.output_path}')
        worker = threading.Thread(
            target=self._run_compare_worker, args=(request,), daemon=True
        )
        worker.start()

    def _mode_name(self, mode: str) -> str:
        if mode == MODE_MULTI:
            return '1 отчёт + много табелей'
        if mode == MODE_BATCH:
            return 'много пар'
        if mode == MODE_EMPLOYEES:
            return 'сверка сотрудников'
        return '1 отчёт + 1 табель'

    def _run_compare_worker(self, request: CompareRequest) -> None:
        try:
            if request.mode == MODE_BATCH:
                saved_count, leftovers = self._execute_batch_request(request)
                message = f'{request.output_path} | файлов: {saved_count}'
                if leftovers:
                    message += ' | без пары: ' + ', '.join(leftovers)
                self._worker_queue.put(('success', message))
            else:
                result_bytes = self._execute_request(request)
                request.output_path.parent.mkdir(parents=True, exist_ok=True)
                request.output_path.write_bytes(result_bytes)
                self._worker_queue.put(('success', str(request.output_path)))
        except Exception as exc:
            details = ''.join(
                traceback.format_exception_only(type(exc), exc)
            ).strip()
            self._worker_queue.put(('error', details))

    def _execute_request(self, request: CompareRequest) -> bytes:
        try:
            if request.mode == MODE_EMPLOYEES:
                assert request.secondary_path is not None
                result = compare_employee_presence(
                    request.primary_path.read_bytes(),
                    request.secondary_path.read_bytes(),
                )
                return _ensure_result_note(result, mode=request.mode)

            config = _build_config(request.period)

            if request.mode == MODE_MULTI:
                result = compare_report_and_many_timesheets(
                    request.primary_path.read_bytes(),
                    [p.read_bytes() for p in request.timesheet_paths],
                    config=config,
                )
                return _ensure_result_note(result, mode=request.mode)

            assert request.secondary_path is not None
            result, _ = _compare_two_excel_files(
                request.primary_path, request.secondary_path, config=config
            )
            return _ensure_result_note(result, mode=request.mode)

        except InputFormatError:
            assert request.secondary_path is not None
            result = compare_employee_presence(
                request.secondary_path.read_bytes(),
                request.primary_path.read_bytes(),
            )
            return _ensure_result_note(result, mode=request.mode)

    def _execute_batch_request(self, request: CompareRequest) -> tuple[int, list[str]]:
        config = _build_config(request.period)
        pending = list(request.batch_paths)
        saved_count = 0
        leftovers: list[str] = []

        request.output_path.mkdir(parents=True, exist_ok=True)

        while len(pending) >= 2:
            best_candidate = None

            for i in range(len(pending) - 1):
                for j in range(i + 1, len(pending)):
                    file_a = pending[i]
                    file_b = pending[j]
                    try:
                        result_bytes, out_name = _compare_two_excel_files(
                            file_a, file_b, config=config
                        )
                    except InputFormatError:
                        continue
                    data_rows = _result_data_row_count(result_bytes)
                    name_score = filename_match_score(file_a.name, file_b.name)
                    has_data = 1 if data_rows > 0 else 0
                    rank = (name_score, has_data, data_rows)
                    if best_candidate is None or rank > tuple(best_candidate[:3]):
                        best_candidate = list(rank) + [i, j, result_bytes, out_name]

            if best_candidate is None:
                break

            name_score, has_data, data_rows, i, j, result_bytes, out_name = best_candidate

            if has_data == 0 and name_score == 0:
                break

            result_bytes = _ensure_result_note(result_bytes, mode=MODE_SINGLE)
            (request.output_path / out_name).write_bytes(result_bytes)
            pending.pop(j)
            pending.pop(i)
            saved_count += 1

        leftovers = [path.name for path in pending]

        if saved_count == 0:
            raise InputFormatError(
                'Не удалось найти ни одной пары отчёт/табель среди выбранных файлов.'
            )

        return saved_count, leftovers

    def _poll_worker_queue(self) -> None:
        try:
            while True:
                kind, payload = self._worker_queue.get_nowait()
                self._busy = False
                if kind == 'success':
                    self.status_var.set(f'Готово: {payload}')
                    self._log(f'Сверка сохранена: {payload}')
                    messagebox.showinfo('Готово', f'Сверка сохранена:\n{payload}')
                else:
                    self.status_var.set('Ошибка при обработке.')
                    self._log(f'Ошибка: {payload}')
                    messagebox.showerror('Ошибка', payload)
        except queue.Empty:
            pass
        self.after(150, self._poll_worker_queue)


def main() -> None:
    app = SverkaApp()
    app.mainloop()


if __name__ == '__main__':
    main()
