from __future__ import annotations
import datetime as dt
import io
import math
import re
import calendar
from dataclasses import dataclass
from typing import Any, Iterable
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


class InputFormatError(ValueError):
    pass


@dataclass
class ComparisonConfig:
    threshold_minutes: int = 19
    lunch_break_minutes: int = 60
    period_start_day: int | None = None
    period_end_day: int | None = None


@dataclass
class _ReportDay:
    fio_key: str
    fio: str
    date: dt.date
    arrival_min: int | None
    departure_min: int | None
    mark_type: str
    report_minutes: int | None


@dataclass
class _ReportShift:
    fio_key: str
    fio: str
    start_dt: dt.datetime
    end_dt: dt.datetime
    report_minutes: int


_PARENS_RE = re.compile(r'\([^)]*\)')
_LETTER_RE = re.compile(r'[A-Za-zА-Яа-яЁё]')

_NOT_PERSON_FIRST_WORDS: set = set() | frozenset({'участок', 'бригада', 'отдел', 'перетарщик', 'бригадир', 'итого', 'цех', 'группа'})

_DATE_FORMATS: tuple[str, ...] = (
    '%d.%m.%Y', '%d.%m.%y', '%d/%m/%Y', '%d/%m/%y', '%d-%m-%Y', '%d-%m-%y',
    '%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d',
    '%d.%m.%Y %H:%M', '%d.%m.%Y %H:%M:%S',
    '%Y-%m-%d %H:%M', '%Y-%m-%d %H:%M:%S',
    '%Y-%m-%dT%H:%M', '%Y-%m-%dT%H:%M:%S',
)

_DATE_IN_TEXT_RE = re.compile(r'(?P<d>\d{1,2}[./-]\d{1,2}[./-]\d{2,4})')
_ISO_DATE_IN_TEXT_RE = re.compile(r'(?P<d>\d{4}[./-]\d{1,2}[./-]\d{1,2})')
_TIMESHEET_GROUP_RE = re.compile(r'^\d{1,2}\.\d{2,3}[A-Za-zА-Яа-я]?\b')

_FIO_PATTERNS = [re.compile(p) for p in [r'\bфио\b', r'сотруд', r'работник']]
_DATE_PATTERNS = [re.compile(p) for p in [r'\bдата\b', r'\bдень\b', r'\bdate\b']]
_REPORT_ARRIVAL_PATTERNS = [re.compile(p) for p in ['приход', r'\bвход\b', 'начал']]
_REPORT_DEPARTURE_PATTERNS = [re.compile(p) for p in ['уход', r'\bвыход\b', 'конец', 'оконч']]
_TIME_PATTERNS = [re.compile(p) for p in ['врем', r'\btime\b']]
_EVENT_PATTERNS = [re.compile(p) for p in ['событ', r'\bevent\b']]
_TIMESHEET_HOURS_PATTERNS = [re.compile(p) for p in [r'\bитог\b', r'\bитого\b', 'час', 'время', 'отработ']]

_RU_WEEKDAY: dict[str, int] = {'пн': 0, 'вт': 1, 'ср': 2, 'чт': 3, 'пт': 4, 'сб': 5, 'вс': 6}


def _extract_fio_tokens(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, float) and math.isnan(value):
        return []
    text = str(value).strip()
    if not text:
        return []
    text = _PARENS_RE.sub(' ', text)
    text = text.replace(',', ' ').replace('.', ' ')
    text = re.sub(r'\s+', ' ', text).strip()
    if not text:
        return []
    tokens = []
    for token in text.split(' '):
        clean = token.strip().strip('-')
        if not clean:
            continue
        if not _LETTER_RE.search(clean):
            continue
        if any(ch.isdigit() for ch in clean):
            continue
        tokens.append(clean)
    return tokens


def _is_person_fio_tokens(tokens: list[str]) -> bool:
    extra_not_person_first_words = set() | frozenset({
        'фасовщик', 'главная', 'уборщица', 'отгрузка', 'уборщик', 'мойщик',
        'приемщик', 'лифтер', 'комплектовщик', 'мастер', 'оператор', 'приемка',
        'кладовщик', 'грузчик', 'работник', 'мойщица', 'коренщик',
    })
    if len(tokens) < 2:
        return False
    if len(tokens[0]) < 2 or len(tokens[1]) < 2:
        return False
    first = tokens[0].lower()
    if first in _NOT_PERSON_FIRST_WORDS or first in extra_not_person_first_words:
        return False
    return True


def normalize_fio(value: Any) -> str:
    tokens = _extract_fio_tokens(value)
    if not _is_person_fio_tokens(tokens):
        return ''
    return ' '.join(token.lower() for token in tokens)


def normalize_fio_for_output(value: Any) -> str:
    tokens = _extract_fio_tokens(value)
    if not _is_person_fio_tokens(tokens):
        return ''
    return ' '.join(tokens)


def fio_for_output(fio_key: str) -> str:
    return fio_key.title()


def parse_date(value: Any) -> dt.date | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        num = float(value)
        if math.isnan(num):
            return None
        if num >= 20000:
            base = dt.datetime(1899, 12, 30)
            return (base + dt.timedelta(days=num)).date()
        return None
    text = str(value).strip()
    if not text or text.lower() == 'nan':
        return None
    m = _DATE_IN_TEXT_RE.search(text) or _ISO_DATE_IN_TEXT_RE.search(text)
    if m:
        candidate = m.group('d')
        for fmt in _DATE_FORMATS:
            try:
                return dt.datetime.strptime(candidate, fmt).date()
            except ValueError:
                continue
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    cleaned = text.replace('Z', '').strip()
    try:
        return dt.datetime.fromisoformat(cleaned).date()
    except ValueError:
        return None


def _parse_time_minutes(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, dt.time):
        return value.hour * 60 + value.minute
    if isinstance(value, dt.datetime):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        num = float(value)
        if math.isnan(num):
            return None
        if 0 <= num < 1:
            return int(round(num * 24 * 60))
        if 1 <= num < 60000:
            frac = num % 1
            return int(round(frac * 24 * 60))
    text = str(value).strip()
    if not text or text.lower() == 'nan':
        return None
    m = re.match(r'^(?P<h>\d{1,2})\s*:\s*(?P<m>\d{1,2})(?:\s*:\s*\d{1,2})?$', text)
    if m:
        return int(m.group('h')) * 60 + int(m.group('m'))
    for fmt in ('%d.%m.%Y %H:%M', '%d.%m.%Y %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d %H:%M:%S', '%Y-%m-%dT%H:%M', '%Y-%m-%dT%H:%M:%S'):
        try:
            parsed = dt.datetime.strptime(text, fmt)
            return parsed.hour * 60 + parsed.minute
        except ValueError:
            continue
    try:
        parsed = dt.datetime.fromisoformat(text.replace('Z', ''))
        return parsed.hour * 60 + parsed.minute
    except ValueError:
        return None


def _parse_duration_minutes(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, dt.timedelta):
        return int(round(value.total_seconds() / 60))
    if isinstance(value, dt.time):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        num = float(value)
        if 0 <= num < 1:
            return int(round(num * 24 * 60))
        return int(round(num * 60))
    text = str(value).strip()
    if not text or text.lower() == 'nan':
        return None
    text = text.replace(',', '.')
    m = re.match(r'^(?P<h>\d+)\s*:\s*(?P<m>\d{1,2})$', text)
    if m:
        return int(m.group('h')) * 60 + int(m.group('m'))
    try:
        return int(round(float(text) * 60))
    except ValueError:
        pass
    m = re.search(r'(?P<h>\d+)\s*ч', text, flags=re.IGNORECASE)
    if m:
        hours = int(m.group('h'))
        minutes = 0
        m2 = re.search(r'(?P<m>\d+)\s*м', text, flags=re.IGNORECASE)
        if m2:
            minutes = int(m2.group('m'))
        return hours * 60 + minutes
    return None


def _minutes_to_hhmm(minutes: int | None) -> str:
    if minutes is None:
        return ''
    total = int(round(float(minutes)))
    total = max(total, 0)
    return f'{total // 60:02d}:{total % 60:02d}'


def _format_date(value: dt.date | None) -> str:
    if not value:
        return ''
    return value.strftime('%d.%m.%Y')


def _format_date_range(start: dt.date, end: dt.date) -> str:
    return f'{_format_date(start)}-{_format_date(end)}'


def _is_in_period(day: dt.date, config: ComparisonConfig) -> bool:
    if config.period_start_day is not None and day.day < int(config.period_start_day):
        return False
    if config.period_end_day is not None and day.day > int(config.period_end_day):
        return False
    return True


_RU_MONTH_SHORT = {
    'янв': 1, 'фев': 2, 'мар': 3, 'апр': 4, 'май': 5, 'июн': 6,
    'июл': 7, 'авг': 8, 'сен': 9, 'окт': 10, 'ноя': 11, 'дек': 12,
}


def _hhmm_to_minutes(val: Any) -> int | None:
    """Convert integer HHMM (e.g. 725 → 07:25 → 445 min) to minutes from midnight."""
    if val is None:
        return None
    try:
        v = int(val)
        if v < 0:
            return None
        h = v // 100
        m = v % 100
        if h > 23 or m > 59:
            return None
        return h * 60 + m
    except (TypeError, ValueError):
        return None


def _parse_notebook_day_header(val: Any) -> tuple[int, int] | None:
    """Parse '1авг.' → (day=1, month=8). Returns (day, month) or None."""
    if not isinstance(val, str):
        return None
    s = val.strip().lower().rstrip('.')
    for abbr, month in _RU_MONTH_SHORT.items():
        if s.endswith(abbr):
            try:
                day = int(s[: -len(abbr)].strip())
                if 1 <= day <= 31:
                    return (day, month)
            except ValueError:
                pass
    return None


def _read_notebook_data(
    notebook_excel_bytes: bytes,
    reference_dates: list[dt.date] | None = None,
) -> tuple[dict, dict]:
    """Parse тетрадь (manual notebook) with HHMM arrival/departure columns per day.

    Returns (minutes_by_day, fio_by_key).
    minutes_by_day: {(fio_key, date): effective_minutes}
      A day with only arrival or only departure is recorded as 0 minutes
      (person was present, just incomplete data).
    """
    wb = _load_workbook_from_bytes(notebook_excel_bytes)
    ws = next((s for s in wb.worksheets if s.max_row > 1), None)
    if ws is None:
        return {}, {}

    year = (reference_dates[0].year if reference_dates else dt.date.today().year)

    header_row: tuple = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

    fio_idx = next(
        (i for i, h in enumerate(header_row) if h and 'фио' in str(h).lower()), 0
    )

    # Build list of (arrival_col, departure_col, date)
    day_columns: list[tuple[int, int, dt.date]] = []
    for i, h in enumerate(header_row):
        parsed = _parse_notebook_day_header(h)
        if parsed is not None:
            day_num, month_num = parsed
            dep_idx = i + 1
            try:
                d = dt.date(year, month_num, day_num)
                day_columns.append((i, dep_idx, d))
            except ValueError:
                pass

    minutes_by_day: dict = {}
    fio_by_key: dict = {}
    default_config = ComparisonConfig()

    for row in ws.iter_rows(min_row=2, values_only=True):
        fio_val = row[fio_idx] if fio_idx < len(row) else None
        if not fio_val:
            continue
        fio_key = normalize_fio(fio_val)
        if not fio_key or not _is_person_fio_tokens(_extract_fio_tokens(fio_val)):
            continue
        fio_by_key.setdefault(fio_key, normalize_fio_for_output(fio_val))

        for arr_idx, dep_idx, day in day_columns:
            arr_raw = row[arr_idx] if arr_idx < len(row) else None
            dep_raw = row[dep_idx] if dep_idx < len(row) else None
            arr_min = _hhmm_to_minutes(arr_raw)
            dep_min = _hhmm_to_minutes(dep_raw)

            if arr_min is None and dep_min is None:
                continue  # no data for this day

            if arr_min is not None and dep_min is not None:
                effective = _compute_report_minutes(arr_min, dep_min, default_config)
            else:
                effective = 0  # present but incomplete record

            key = (fio_key, day)
            minutes_by_day[key] = minutes_by_day.get(key, 0) + effective

    return minutes_by_day, fio_by_key


def _normalize_header(value: Any) -> str:
    text = str(value).strip().lower()
    text = re.sub(r'\s+', ' ', text)
    return text


def _find_column_index(headers: Iterable[Any], patterns: list[re.Pattern[str]]) -> int | None:
    for idx, col in enumerate(headers):
        header = _normalize_header(col)
        for p in patterns:
            if p.search(header):
                return idx
    return None


def _load_workbook_from_bytes(excel_bytes: bytes):
    try:
        return load_workbook(io.BytesIO(excel_bytes), data_only=True)
    except Exception as e:
        raise InputFormatError('Не удалось прочитать Excel. Подходит формат .xlsx.') from e


def _iter_header_candidates(ws: Worksheet, max_header_rows: int = 20):
    last = min(max_header_rows, ws.max_row)
    for header_row in range(1, last + 1):
        rows = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        if not rows:
            continue
        headers = list(rows[0])
        if any(str(h).strip() if h is not None else False for h in headers):
            yield header_row, headers


def _get_cell(row: tuple[Any, ...], idx: int) -> Any:
    if idx < len(row):
        return row[idx]
    return None


def _select_report_layout(wb) -> tuple:
    last_error = None
    for ws in wb.worksheets:
        for header_row, headers in _iter_header_candidates(ws):
            fio = _find_column_index(headers, _FIO_PATTERNS)
            date = _find_column_index(headers, _DATE_PATTERNS)
            arrival = _find_column_index(headers, _REPORT_ARRIVAL_PATTERNS)
            departure = _find_column_index(headers, _REPORT_DEPARTURE_PATTERNS)
            if fio is not None and date is not None and arrival is not None and departure is not None:
                return ws, header_row, fio, date, arrival, departure
            if fio is None and date is None:
                continue
            missing = []
            if fio is None:
                missing.append('ФИО')
            if date is None:
                missing.append('Дата')
            if arrival is None:
                missing.append('Приход')
            if departure is None:
                missing.append('Уход')
            headers_text = ', '.join(str(h) for h in headers if h is not None)
            last_error = InputFormatError(
                f'В отчёте не нашёл колонки: {", ".join(missing)}. Заголовки: {headers_text}'
            )
    if last_error:
        raise last_error
    raise InputFormatError('В отчёте не нашёл нужные колонки: ФИО, Дата, Приход, Уход.')


def _compute_report_minutes(arrival_min: int, departure_min: int, config: ComparisonConfig) -> int:
    diff = departure_min - arrival_min
    if diff < 0:
        diff += 1440
    # Обед вычитается если смена >= 6 часов (360 минут)
    if diff >= 360:
        diff -= config.lunch_break_minutes
    return max(diff, 0)


def _minutes_to_time(minutes: int) -> dt.time:
    minutes = max(0, int(minutes))
    minutes = minutes % 1440
    return dt.time(minutes // 60, minutes % 60)


def _read_report_shifts_event_log(wb, config: ComparisonConfig) -> tuple[list[_ReportShift], list[_ReportDay]]:
    last_error = None
    for ws in wb.worksheets:
        for header_row, headers in _iter_header_candidates(ws, max_header_rows=40):
            fio_idx = _find_column_index(headers, _FIO_PATTERNS)
            date_idx = _find_column_index(headers, _DATE_PATTERNS)
            time_idx = _find_column_index(headers, _TIME_PATTERNS)
            event_idx = _find_column_index(headers, _EVENT_PATTERNS)

            if fio_idx is None or date_idx is None or time_idx is None or event_idx is None:
                continue

            events: dict[str, list] = {}
            fio_display_by_key: dict[str, str] = {}

            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                fio_raw = _get_cell(row, fio_idx)
                fio_key = normalize_fio(fio_raw)
                if not fio_key:
                    continue
                fio_display = normalize_fio_for_output(fio_raw)
                if fio_display:
                    fio_display_by_key.setdefault(fio_key, fio_display)
                day = parse_date(_get_cell(row, date_idx))
                if day is None:
                    continue
                t_min = _parse_time_minutes(_get_cell(row, time_idx))
                if t_min is None:
                    continue
                event_raw = _get_cell(row, event_idx)
                event_text = str(event_raw).strip().lower() if event_raw is not None else ''
                if not event_text:
                    continue
                if 'вход' in event_text:
                    et = 'in'
                elif 'выход' in event_text:
                    et = 'out'
                else:
                    continue
                ts = dt.datetime.combine(day, _minutes_to_time(t_min))
                events.setdefault(fio_key, []).append((ts, et))

            shifts: list[_ReportShift] = []
            one_sided: list[_ReportDay] = []

            # Максимальный перерыв внутри одной смены (4 часа).
            # Если человек отсутствовал дольше — это конец смены, не перерыв.
            _MAX_BREAK_MIN = 240

            def _record_shift(fio_key, fio, shift_start, shift_end, mid_gap_min):
                total_raw = int((shift_end - shift_start).total_seconds() // 60)
                if mid_gap_min > 0:
                    # Человек выходил в течение смены
                    if total_raw >= 360:
                        # Если перерыв < часа — вычитаем стандартный обед (60 мин),
                        # перерыв поглощается обедом. Если >= часа — вычитаем фактический
                        # перерыв целиком (он уже включает стандартный обед + сверх).
                        deduction = max(mid_gap_min, config.lunch_break_minutes)
                    else:
                        # Смена < 6 часов — обед не предусмотрен, вычитаем фактическое время выхода
                        deduction = mid_gap_min
                else:
                    # Выходов внутри смены не было — стандартное вычитание обеда
                    deduction = config.lunch_break_minutes if total_raw >= 360 else 0
                effective = max(total_raw - deduction, 0)
                shifts.append(_ReportShift(
                    fio_key=fio_key, fio=fio,
                    start_dt=shift_start, end_dt=shift_end,
                    report_minutes=effective,
                ))

            for fio_key, seq in events.items():
                seq.sort(key=lambda x: x[0])
                fio = fio_display_by_key.get(fio_key, fio_for_output(fio_key))
                open_in: dt.datetime | None = None
                gap_start: dt.datetime | None = None  # момент выхода в перерыве
                mid_gap: int = 0  # суммарный перерыв внутри текущей смены (мин)

                for ts, et in seq:
                    if et == 'in':
                        if open_in is None:
                            # Начало новой смены
                            open_in = ts
                            mid_gap = 0
                            gap_start = None
                        elif gap_start is not None:
                            # Вернулся с перерыва
                            gap_dur = int((ts - gap_start).total_seconds() // 60)
                            if gap_dur > _MAX_BREAK_MIN:
                                # Слишком длинный перерыв — закрываем предыдущую смену
                                _record_shift(fio_key, fio, open_in, gap_start, mid_gap)
                                open_in = ts
                                mid_gap = 0
                            else:
                                mid_gap += gap_dur
                            gap_start = None
                        # else: дублирующий вход пока уже внутри смены без gap — игнорируем
                    else:  # et == 'out'
                        if open_in is None:
                            # Выход без входа — односторонняя отметка
                            one_sided.append(_ReportDay(
                                fio_key=fio_key, fio=fio, date=ts.date(),
                                arrival_min=None,
                                departure_min=ts.hour * 60 + ts.minute,
                                mark_type='departure_only', report_minutes=None,
                            ))
                        elif gap_start is not None:
                            # Дублирующий выход, игнорируем
                            pass
                        else:
                            span = int((ts - open_in).total_seconds() // 60)
                            if span > 1620:
                                # Вход висит больше 27 часов — односторонняя отметка на вход
                                one_sided.append(_ReportDay(
                                    fio_key=fio_key, fio=fio, date=open_in.date(),
                                    arrival_min=open_in.hour * 60 + open_in.minute,
                                    departure_min=None,
                                    mark_type='arrival_only', report_minutes=None,
                                ))
                                open_in = None
                                # Сам выход теперь висит без входа — тоже односторонний
                                one_sided.append(_ReportDay(
                                    fio_key=fio_key, fio=fio, date=ts.date(),
                                    arrival_min=None,
                                    departure_min=ts.hour * 60 + ts.minute,
                                    mark_type='departure_only', report_minutes=None,
                                ))
                            else:
                                # Выход — либо конец смены, либо начало перерыва;
                                # узнаем по следующему событию (если вход вернётся — перерыв)
                                gap_start = ts

                # Конец событий для этого сотрудника
                if gap_start is not None and open_in is not None:
                    # Последнее событие было выход — это конец смены
                    span = int((gap_start - open_in).total_seconds() // 60)
                    if span > 1620:
                        one_sided.append(_ReportDay(
                            fio_key=fio_key, fio=fio, date=open_in.date(),
                            arrival_min=open_in.hour * 60 + open_in.minute,
                            departure_min=None,
                            mark_type='arrival_only', report_minutes=None,
                        ))
                    else:
                        _record_shift(fio_key, fio, open_in, gap_start, mid_gap)
                    open_in = None
                if open_in is not None:
                    # Незакрытый вход — односторонняя отметка
                    one_sided.append(_ReportDay(
                        fio_key=fio_key, fio=fio, date=open_in.date(),
                        arrival_min=open_in.hour * 60 + open_in.minute,
                        departure_min=None,
                        mark_type='arrival_only', report_minutes=None,
                    ))

            if shifts:
                return shifts, one_sided
            last_error = InputFormatError('В отчёте (лог событий) не нашёл ни одной пары Вход/Выход.')
    if last_error:
        raise last_error
    raise InputFormatError('Не смог распознать отчёт: нужен формат с колонками Дата/Время/Сотрудник/Событие.')


def _report_days_to_shifts(report_days: list[_ReportDay], config: ComparisonConfig) -> tuple[list[_ReportShift], list[_ReportDay]]:
    """
    Возвращает (shifts, one_sided_days).
    shifts — полные смены (есть и приход, и уход).
    one_sided_days — односторонние отметки (только приход или только уход).
    """
    shifts: list[_ReportShift] = []
    one_sided_days: list[_ReportDay] = []

    for d in report_days:
        # ИСПРАВЛЕНИЕ 2: Односторонние отметки — собираем отдельно
        if d.mark_type == 'arrival_only' or d.mark_type == 'departure_only':
            one_sided_days.append(d)
            continue
        if d.mark_type != 'both':
            continue
        if d.arrival_min is None or d.departure_min is None:
            continue

        start_dt = dt.datetime.combine(d.date, _minutes_to_time(d.arrival_min))
        end_dt = dt.datetime.combine(d.date, _minutes_to_time(d.departure_min))
        if d.departure_min < d.arrival_min:
            end_dt += dt.timedelta(days=1)

        if d.report_minutes is not None:
            minutes = d.report_minutes
        else:
            minutes = _compute_report_minutes(int(d.arrival_min), int(d.departure_min), config)

        shifts.append(_ReportShift(
            fio_key=d.fio_key,
            fio=d.fio,
            start_dt=start_dt,
            end_dt=end_dt,
            report_minutes=int(minutes),
        ))

    return shifts, one_sided_days


def _read_report_shifts(report_excel_bytes: bytes, config: ComparisonConfig) -> tuple[list[_ReportShift], list[_ReportDay]]:
    wb = _load_workbook_from_bytes(report_excel_bytes)
    try:
        shifts, one_sided_days = _read_report_shifts_event_log(wb, config)
        return shifts, one_sided_days
    except InputFormatError:
        report_days = _read_report_days(report_excel_bytes, config)
        shifts, one_sided_days = _report_days_to_shifts(report_days, config)
        if shifts:
            return shifts, one_sided_days
        raise


def _read_report_employees_and_dates(report_excel_bytes: bytes) -> tuple[dict[str, str], list[dt.date]]:
    wb = _load_workbook_from_bytes(report_excel_bytes)
    out: dict[str, str] = {}
    dates: set[dt.date] = set()

    for ws in wb.worksheets:
        for header_row, headers in _iter_header_candidates(ws, max_header_rows=40):
            fio_idx = _find_column_index(headers, _FIO_PATTERNS)
            date_idx = _find_column_index(headers, _DATE_PATTERNS)
            time_idx = _find_column_index(headers, _TIME_PATTERNS)
            event_idx = _find_column_index(headers, _EVENT_PATTERNS)

            if fio_idx is None or date_idx is None or time_idx is None or event_idx is None:
                continue

            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                fio_raw = _get_cell(row, fio_idx)
                fio_key = normalize_fio(fio_raw)
                if not fio_key:
                    continue
                fio_out = normalize_fio_for_output(fio_raw)
                if fio_out:
                    out.setdefault(fio_key, fio_out)
                day = parse_date(_get_cell(row, date_idx))
                if day:
                    dates.add(day)

            if out:
                return out, sorted(dates)

    if out:
        return out, sorted(dates)
    if not out:
        raise InputFormatError('Не смог распознать отчёт для сверки сотрудников.')
    raise InputFormatError('В отчёте не нашёл сотрудников.')


def _build_report_days_from_acc(acc: dict, config: ComparisonConfig) -> list[_ReportDay]:
    days: list[_ReportDay] = []
    for (fio_key, date), rec in acc.items():
        arrival_min = rec.get('arrival_min')
        departure_min = rec.get('departure_min')

        if arrival_min is not None and departure_min is not None:
            mark_type = 'both'
            report_minutes = _compute_report_minutes(int(arrival_min), int(departure_min), config)
        elif arrival_min is not None and departure_min is None:
            mark_type = 'arrival_only'
            report_minutes = None
        elif arrival_min is None and departure_min is not None:
            mark_type = 'departure_only'
            report_minutes = None
        else:
            mark_type = 'none'
            report_minutes = None

        fio = str(rec.get('fio', '') or '') or fio_for_output(fio_key)

        days.append(_ReportDay(
            fio_key=fio_key,
            fio=fio,
            date=date,
            arrival_min=arrival_min,
            departure_min=departure_min,
            mark_type=mark_type,
            report_minutes=report_minutes,
        ))
    return days


def _is_blank_cell(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if isinstance(value, str):
        return not value.strip()
    return False


def _date_header_from_row(row: tuple[Any, ...]) -> dt.date | None:
    if not row:
        return None
    d = parse_date(row[0])
    if d is None:
        return None
    if any(not _is_blank_cell(v) for v in row[1:]):
        return None
    return d


def _read_report_days_table(wb, config: ComparisonConfig) -> list[_ReportDay]:
    ws, header_row, fio_idx, date_idx, arrival_idx, departure_idx = _select_report_layout(wb)
    acc: dict = {}

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        fio_key = normalize_fio(_get_cell(row, fio_idx))
        if not fio_key:
            continue
        date = parse_date(_get_cell(row, date_idx))
        if date is None:
            continue
        arrival_min = _parse_time_minutes(_get_cell(row, arrival_idx))
        departure_min = _parse_time_minutes(_get_cell(row, departure_idx))
        key = (fio_key, date)
        rec = acc.get(key)
        if rec is None:
            fio_out = normalize_fio_for_output(_get_cell(row, fio_idx)) or fio_for_output(fio_key)
            rec = {'fio': fio_out, 'arrival_min': None, 'departure_min': None}
            acc[key] = rec
        if arrival_min is not None:
            rec['arrival_min'] = arrival_min
        if departure_min is not None:
            rec['departure_min'] = departure_min

    return _build_report_days_from_acc(acc, config)


def _read_report_days_by_date_blocks(wb, config: ComparisonConfig) -> list[_ReportDay]:
    acc: dict = {}
    found_date_headers = 0
    found_employee_rows = 0

    for ws in wb.worksheets:
        current_date = None
        arrival_idx = None
        departure_idx = None

        for row in ws.iter_rows(values_only=True):
            date_header = _date_header_from_row(row)
            if date_header is not None:
                current_date = date_header
                arrival_idx = None
                departure_idx = None
                found_date_headers += 1
                continue

            if current_date is None:
                continue

            if arrival_idx is None or departure_idx is None:
                a = _find_column_index(row, _REPORT_ARRIVAL_PATTERNS)
                d = _find_column_index(row, _REPORT_DEPARTURE_PATTERNS)
                if a is not None and d is not None and a != d:
                    arrival_idx = a
                    departure_idx = d
                    continue

            fio_key = normalize_fio(_get_cell(row, 0))
            if not fio_key:
                continue

            a_idx = arrival_idx if arrival_idx is not None else 1
            d_idx = departure_idx if departure_idx is not None else 2

            arrival_min = _parse_time_minutes(_get_cell(row, a_idx))
            departure_min = _parse_time_minutes(_get_cell(row, d_idx))

            key = (fio_key, current_date)
            if key not in acc:
                fio_out = normalize_fio_for_output(_get_cell(row, 0)) or fio_for_output(fio_key)
                acc[key] = {'fio': fio_out, 'arrival_min': None, 'departure_min': None}
            if arrival_min is not None:
                acc[key]['arrival_min'] = arrival_min
            if departure_min is not None:
                acc[key]['departure_min'] = departure_min
            found_employee_rows += 1

    return _build_report_days_from_acc(acc, config)


def _read_report_days(report_excel_bytes: bytes, config: ComparisonConfig) -> list[_ReportDay]:
    wb = _load_workbook_from_bytes(report_excel_bytes)
    try:
        days = _read_report_days_table(wb, config)
        if days:
            return days
    except InputFormatError as table_error:
        days = _read_report_days_by_date_blocks(wb, config)
        if days:
            return days
        raise table_error

    days = _read_report_days_by_date_blocks(wb, config)
    if days:
        return days
    raise InputFormatError('В отчёте не нашёл данные по датам и сотрудникам.')


def _parse_timesheet_header_date(value: Any, reference_dates: list[dt.date] | None) -> dt.date | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    d = parse_date(value)
    if d is not None:
        if reference_dates is None or d in reference_dates:
            return d
        return None
    m = re.match(r'^(?P<day>\d{1,2})(?:\s+|$)(?P<wd>[а-я]{2})?', str(value).strip().lower())
    if m:
        try:
            day_num = int(m.group('day'))
            wd = m.group('wd') or ''
            if reference_dates:
                # When weekday is present, use it to disambiguate between months
                for ref in reference_dates:
                    if ref.day == day_num:
                        if wd and wd in _RU_WEEKDAY:
                            if ref.weekday() == _RU_WEEKDAY[wd]:
                                return ref
                        else:
                            return ref
        except (ValueError, TypeError):
            pass
    return None


def _parse_timesheet_header_day_weekday(value: Any) -> tuple[int, str] | None:
    if value is None:
        return None
    text = str(value).strip().lower()
    m = re.match(r'^(?P<day>\d{1,2})(?:\s+|$)(?P<wd>[а-я]{2})?', text)
    if not m:
        return None
    try:
        day_num = int(m.group('day'))
        wd = m.group('wd') or ''
        return day_num, wd
    except (ValueError, TypeError):
        return None


def _infer_timesheet_month_date_columns(headers: list[Any], fio_idx: int, reference_dates: list[dt.date] | None) -> list[tuple[int, dt.date]]:
    if not reference_dates:
        return []
    month_candidates = sorted({(d.year, d.month) for d in reference_dates})
    if not month_candidates:
        return []

    parsed_headers = []
    for idx, h in enumerate(headers):
        if idx == fio_idx:
            continue
        parsed = _parse_timesheet_header_day_weekday(h)
        if parsed is None:
            continue
        parsed_headers.append((idx, parsed[0], parsed[1]))

    if not parsed_headers:
        return []

    best_month = None
    best_score = -1
    for year, month in month_candidates:
        score = 0
        mdays = calendar.monthrange(year, month)[1]
        for _, day_num, wd in parsed_headers:
            if 1 <= day_num <= mdays:
                try:
                    actual_date = dt.date(year, month, day_num)
                    if wd and wd in _RU_WEEKDAY:
                        if actual_date.weekday() == _RU_WEEKDAY[wd]:
                            score += 3
                        else:
                            score -= 1
                    else:
                        score += 1
                except ValueError:
                    pass
        if score > best_score:
            best_score = score
            best_month = (year, month)

    if best_month is None:
        return []

    out = []
    year, month = best_month
    mdays = calendar.monthrange(year, month)[1]
    for idx, day_num, _ in parsed_headers:
        if 1 <= day_num <= mdays:
            try:
                day = dt.date(year, month, day_num)
                out.append((idx, day))
            except ValueError:
                pass
    return out


def _is_timesheet_non_employee_row(row: tuple, fio_idx: int) -> bool:
    if fio_idx >= len(row):
        return False
    fio_cell = row[fio_idx]
    fio_value = getattr(fio_cell, 'value', fio_cell)
    if fio_value is None:
        return False
    text = str(fio_value).strip()
    if not text:
        return False
    lowered = text.lower()
    if lowered.startswith('итого'):
        return True
    if _TIMESHEET_GROUP_RE.match(text):
        return True
    is_bold = bool(getattr(getattr(fio_cell, 'font', None), 'bold', False))
    if is_bold:
        return True
    # '#' in column 0 only indicates a section header when the FIO column is also 0
    if fio_idx == 0:
        first_text = str(fio_value).strip()
        if first_text.startswith('#'):
            return True
    return False


def _build_timesheet_long(ws: Worksheet, header_row: int, fio_idx: int, date_idx: int, hours_idx: int) -> tuple[dict, dict]:
    acc: dict = {}
    fio_by_key: dict = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        if _is_timesheet_non_employee_row(row, fio_idx):
            continue
        values = tuple(getattr(c, 'value', c) for c in row)
        fio_raw = _get_cell(values, fio_idx)
        fio_key = normalize_fio(fio_raw)
        if not fio_key:
            continue
        fio_out = normalize_fio_for_output(fio_raw)
        if fio_out:
            fio_by_key.setdefault(fio_key, fio_out)
        date = parse_date(_get_cell(values, date_idx))
        if date is None:
            continue
        minutes = _parse_duration_minutes(_get_cell(values, hours_idx))
        key = (fio_key, date)
        rec = acc.get(key)
        if rec is None:
            rec = {'sum': 0, 'has': False}
            acc[key] = rec
        if minutes is not None:
            rec['sum'] += minutes
            rec['has'] = True
    out = {k: v['sum'] if v['has'] else None for k, v in acc.items()}
    return out, fio_by_key


def _build_timesheet_wide(ws: Worksheet, header_row: int, fio_idx: int, date_columns: list[tuple[int, dt.date]]) -> tuple[dict, dict]:
    out: dict = {}
    fio_by_key: dict = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        if _is_timesheet_non_employee_row(row, fio_idx):
            continue
        values = tuple(getattr(c, 'value', c) for c in row)
        fio_raw = _get_cell(values, fio_idx)
        fio_key = normalize_fio(fio_raw)
        if not fio_key:
            continue
        fio_out = normalize_fio_for_output(fio_raw)
        if fio_out:
            fio_by_key.setdefault(fio_key, fio_out)
        for col_idx, day in date_columns:
            minutes = _parse_duration_minutes(_get_cell(values, col_idx))
            if minutes is None:
                continue
            key = (fio_key, day)
            out[key] = out.get(key, 0) + minutes
    for k, v in list(out.items()):
        if v is not None:
            out[k] = v
    return out, fio_by_key


def _read_timesheet_minutes(timesheet_excel_bytes: bytes, reference_dates: list[dt.date] | None = None) -> tuple[dict, dict]:
    wb = _load_workbook_from_bytes(timesheet_excel_bytes)
    last_error = None

    for ws in wb.worksheets:
        for header_row, headers in _iter_header_candidates(ws):
            fio_idx = _find_column_index(headers, _FIO_PATTERNS)
            if fio_idx is None:
                continue

            date_idx = _find_column_index(headers, _DATE_PATTERNS)
            hours_idx = _find_column_index(headers, _TIMESHEET_HOURS_PATTERNS)

            if date_idx is not None and hours_idx is not None:
                data, fio_by_key = _build_timesheet_long(ws, header_row, fio_idx, date_idx, hours_idx)
                if data:
                    return data, fio_by_key
                last_error = InputFormatError('В табеле (ФИО/Дата/Итог) не нашёл ни одной строки с датой.')
                continue

            date_columns = []
            for idx, h in enumerate(headers):
                if idx == fio_idx:
                    continue
                d = _parse_timesheet_header_date(h, reference_dates)
                if d is not None:
                    date_columns.append((idx, d))

            if not date_columns and reference_dates:
                inferred = _infer_timesheet_month_date_columns(headers, fio_idx, reference_dates)
                date_columns = inferred

            if date_columns:
                data, fio_by_key = _build_timesheet_wide(ws, header_row, fio_idx, date_columns)
                if data:
                    return data, fio_by_key
                last_error = InputFormatError('В табеле (широкий формат) не нашёл ни одной строки с ФИО.')
                continue

    if last_error:
        raise last_error
    raise InputFormatError(
        "Не смог распознать табель: нужна либо колонка 'Дата' + колонка часов (Итог/Часы), "
        "либо табель в широком виде с датами в заголовках."
    )


def _read_timesheet_employees(timesheet_excel_bytes: bytes, reference_dates: list[dt.date] | None = None) -> dict[str, str]:
    wb = _load_workbook_from_bytes(timesheet_excel_bytes)
    last_error = None

    for ws in wb.worksheets:
        for header_row, headers in _iter_header_candidates(ws):
            fio_idx = _find_column_index(headers, _FIO_PATTERNS)
            if fio_idx is None:
                continue

            date_idx = _find_column_index(headers, _DATE_PATTERNS)
            hours_idx = _find_column_index(headers, _TIMESHEET_HOURS_PATTERNS)

            is_long = date_idx is not None and hours_idx is not None
            is_wide = False
            date_columns = []

            if not is_long:
                for idx, h in enumerate(headers):
                    if idx == fio_idx:
                        continue
                    d = _parse_timesheet_header_date(h, reference_dates)
                    if d is not None:
                        date_columns.append((idx, d))
                if not date_columns and reference_dates:
                    date_columns = _infer_timesheet_month_date_columns(headers, fio_idx, reference_dates)
                is_wide = bool(date_columns)

            if not is_long and not is_wide:
                continue

            out: dict[str, str] = {}
            for row in ws.iter_rows(min_row=header_row + 1):
                if _is_timesheet_non_employee_row(row, fio_idx):
                    continue
                values = tuple(getattr(c, 'value', c) for c in row)
                fio_raw = _get_cell(values, fio_idx)
                fio_key = normalize_fio(fio_raw)
                if not fio_key:
                    continue
                fio_out = normalize_fio_for_output(fio_raw) or ''
                if fio_out:
                    out.setdefault(fio_key, fio_out)

            if out:
                return out
            last_error = InputFormatError('В табеле не нашёл сотрудников.')

    if last_error:
        raise last_error
    raise InputFormatError('Не смог распознать табель для сверки сотрудников.')


def _read_employee_list_employees(employee_list_excel_bytes: bytes) -> dict[str, str]:
    wb = _load_workbook_from_bytes(employee_list_excel_bytes)
    last_error = None

    for ws in wb.worksheets:
        for header_row, headers in _iter_header_candidates(ws, max_header_rows=20):
            fio_idx = _find_column_index(headers, _FIO_PATTERNS)
            if fio_idx is None:
                continue

            out: dict[str, str] = {}
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                fio_raw = _get_cell(row, fio_idx)
                fio_key = normalize_fio(fio_raw)
                if not fio_key:
                    continue
                fio_out = normalize_fio_for_output(fio_raw)
                if not fio_out:
                    continue
                out.setdefault(fio_key, fio_out)

            if out:
                return out
            last_error = InputFormatError('В списке сотрудников не нашёл сотрудников.')

    if last_error:
        raise last_error
    raise InputFormatError('Не смог распознать список сотрудников.')


def compare_employee_presence(employee_list_excel_bytes: bytes, timesheet_excel_bytes: bytes) -> bytes:
    employee_list_fio_by_key = _read_employee_list_employees(employee_list_excel_bytes)
    timesheet_fio_by_key = _read_timesheet_employees(timesheet_excel_bytes)

    only_timesheet_keys = sorted(k for k in timesheet_fio_by_key if k not in employee_list_fio_by_key)
    only_employee_list_keys = sorted(k for k in employee_list_fio_by_key if k not in timesheet_fio_by_key)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = 'Только в табеле'
    ws1.append(['ФИО', 'Файл'])
    for key in only_timesheet_keys:
        ws1.append([timesheet_fio_by_key.get(key, fio_for_output(key)), 'Табель'])

    ws2 = wb.create_sheet('Только в списке')
    ws2.append(['ФИО', 'Файл'])
    for key in only_employee_list_keys:
        ws2.append([employee_list_fio_by_key.get(key, fio_for_output(key)), 'Список сотрудников'])

    ws3 = wb.create_sheet('Количество сотрудников')
    ws3.append(['Показатель', 'Количество'])
    ws3.append(['Уникальных сотрудников в табеле', len(timesheet_fio_by_key)])
    ws3.append(['Уникальных сотрудников в списке', len(employee_list_fio_by_key)])
    ws3.append(['Сотрудников в обоих файлах', len(set(timesheet_fio_by_key) & set(employee_list_fio_by_key))])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def compare_report_and_timesheet(
    report_excel_bytes: bytes,
    timesheet_excel_bytes: bytes,
    config: ComparisonConfig | None = None,
    notebook_excel_bytes: bytes | None = None,
) -> bytes:
    if not config:
        config = ComparisonConfig()

    report_shifts, one_sided_days = _read_report_shifts(report_excel_bytes, config)
    if not report_shifts:
        raise InputFormatError('В отчёте не нашёл смен/пар Вход-Выход.')

    min_day = min(s.start_dt.date() for s in report_shifts)
    max_day = max(s.end_dt.date() for s in report_shifts)
    ref_start = dt.date(min_day.year, min_day.month, 1)
    last_day = calendar.monthrange(max_day.year, max_day.month)[1]
    ref_end = dt.date(max_day.year, max_day.month, last_day)
    reference_dates = [ref_start + dt.timedelta(days=i) for i in range((ref_end - ref_start).days + 1)]

    timesheet_minutes, timesheet_fio_by_key = _read_timesheet_minutes(timesheet_excel_bytes, reference_dates=reference_dates)

    notebook_minutes: dict = {}
    notebook_fio_by_key: dict = {}
    if notebook_excel_bytes:
        notebook_minutes, notebook_fio_by_key = _read_notebook_data(notebook_excel_bytes, reference_dates)

    return _build_output_workbook(
        report_shifts, timesheet_minutes, timesheet_fio_by_key, config, one_sided_days,
        notebook_minutes=notebook_minutes, notebook_fio_by_key=notebook_fio_by_key,
    )


def compare_report_and_many_timesheets(
    report_excel_bytes: bytes,
    timesheet_excel_list: list[bytes],
    config: ComparisonConfig | None = None,
) -> bytes:
    if not config:
        config = ComparisonConfig()
    if not timesheet_excel_list:
        raise InputFormatError('Не получил ни одного табеля для сверки.')

    report_shifts, one_sided_days = _read_report_shifts(report_excel_bytes, config)
    if not report_shifts:
        raise InputFormatError('В отчёте не нашёл смен/пар Вход-Выход.')

    min_day = min(s.start_dt.date() for s in report_shifts)
    max_day = max(s.end_dt.date() for s in report_shifts)
    ref_start = dt.date(min_day.year, min_day.month, 1)
    last_day = calendar.monthrange(max_day.year, max_day.month)[1]
    ref_end = dt.date(max_day.year, max_day.month, last_day)
    reference_dates = [ref_start + dt.timedelta(days=i) for i in range((ref_end - ref_start).days + 1)]

    merged_timesheet_minutes: dict = {}
    merged_timesheet_fio_by_key: dict = {}

    for timesheet_bytes in timesheet_excel_list:
        timesheet_minutes, timesheet_fio_by_key = _read_timesheet_minutes(timesheet_bytes, reference_dates=reference_dates)
        for (fio_key, day), minutes in timesheet_minutes.items():
            key = (fio_key, day)
            merged_timesheet_minutes[key] = (merged_timesheet_minutes.get(key) or 0) + (minutes or 0)
        for fio_key, fio_out in timesheet_fio_by_key.items():
            merged_timesheet_fio_by_key.setdefault(fio_key, fio_out)

    return _build_output_workbook(report_shifts, merged_timesheet_minutes, merged_timesheet_fio_by_key, config, one_sided_days)


def _build_output_workbook(
    report_shifts: list[_ReportShift],
    timesheet_minutes: dict,
    timesheet_fio_by_key: dict,
    config: ComparisonConfig,
    one_sided_days: list[_ReportDay] | None = None,
    notebook_minutes: dict | None = None,
    notebook_fio_by_key: dict | None = None,
) -> bytes:
    notebook_minutes = notebook_minutes or {}
    notebook_fio_by_key = notebook_fio_by_key or {}

    # Агрегируем минуты из отчёта по (fio_key, day)
    report_minutes_by_day: dict = {}
    fio_by_key: dict = {}
    for s in report_shifts:
        start_day = s.start_dt.date()
        if not _is_in_period(start_day, config):
            continue
        key = (s.fio_key, start_day)
        report_minutes_by_day[key] = report_minutes_by_day.get(key, 0) + int(s.report_minutes)
        fio_by_key.setdefault(s.fio_key, s.fio)

    # Объединяем отчёт и тетрадь: combined_minutes_by_day используется вместо report_minutes_by_day
    combined_minutes_by_day: dict = dict(report_minutes_by_day)
    for key, nb_min in notebook_minutes.items():
        if not _is_in_period(key[1], config):
            continue
        combined_minutes_by_day[key] = combined_minutes_by_day.get(key, 0) + nb_min

    # Все сотрудники, у которых есть данные в отчёте или тетради
    combined_fio_keys = (
        set(fio_by_key.keys())
        | {d.fio_key for d in (one_sided_days or [])}
        | set(notebook_fio_by_key.keys())
    )

    def _fio_display(fio_key: str) -> str:
        return (
            fio_by_key.get(fio_key)
            or notebook_fio_by_key.get(fio_key)
            or timesheet_fio_by_key.get(fio_key)
            or fio_for_output(fio_key)
        )

    # Лист 1: Табель > отчёта (табель превышает combined более чем на threshold)
    sheet1 = []
    for (fio_key, day), combined_minutes in combined_minutes_by_day.items():
        t_minutes = int(timesheet_minutes.get((fio_key, day)) or 0)
        diff = t_minutes - int(combined_minutes)
        if diff <= config.threshold_minutes:
            continue
        sheet1.append((_fio_display(fio_key), day, t_minutes, int(combined_minutes), diff))

    # Лист 2: В табеле есть, но нет ни в отчёте, ни в тетради
    one_sided_keys = {(d.fio_key, d.date) for d in (one_sided_days or [])}
    sheet2 = []
    for (fio_key, day), t_minutes in timesheet_minutes.items():
        if not _is_in_period(day, config):
            continue
        if t_minutes is None or t_minutes <= 0:
            continue
        if (fio_key, day) in combined_minutes_by_day:
            continue
        if (fio_key, day) in one_sided_keys:
            continue
        # Если сотрудника нет ни в отчёте, ни в тетради совсем — нет скана, пропускаем
        if fio_key not in combined_fio_keys:
            continue
        fio_value = (
            notebook_fio_by_key.get(fio_key)
            or timesheet_fio_by_key.get(fio_key)
            or fio_for_output(fio_key)
        )
        sheet2.append((fio_value, day, 0, int(t_minutes)))

    sheet1.sort(key=lambda r: (r[0], r[1]))
    sheet2.sort(key=lambda r: (r[0], r[1]))

    wb = Workbook()
    ws1 = wb.active
    ws1.title = 'Табель > отчёта'
    ws1.append(['ФИО', 'Дата', 'Время в табеле', 'Время в отчете', 'Итог'])
    for fio, day, t_minutes, report_minutes, diff in sheet1:
        ws1.append([
            fio,
            _format_date(day),
            _minutes_to_hhmm(t_minutes),
            _minutes_to_hhmm(report_minutes),
            _minutes_to_hhmm(diff),
        ])

    ws2 = wb.create_sheet('В табеле есть, в отчёте нет')
    ws2.append(['ФИО', 'Дата', 'Количество часов в отчете', 'Количество часов в табеле'])
    for fio, day, r_minutes, t_minutes in sheet2:
        ws2.append([
            fio,
            _format_date(day),
            _minutes_to_hhmm(r_minutes),
            _minutes_to_hhmm(t_minutes),
        ])

    # ИСПРАВЛЕНИЕ 2: Лист 3 — Односторонние отметки
    # Работник был на работе, но не отметился (есть только вход ИЛИ только выход)
    if one_sided_days:
        ws3 = wb.create_sheet('Односторонние отметки')
        ws3.append(['ФИО', 'Дата', 'Тип отметки', 'Время'])
        one_sided_sorted = sorted(one_sided_days, key=lambda d: (d.fio, d.date))
        for d in one_sided_sorted:
            if not _is_in_period(d.date, config):
                continue
            if d.mark_type == 'arrival_only':
                mark_label = 'Только приход (нет ухода)'
                time_str = _minutes_to_hhmm(d.arrival_min)
            elif d.mark_type == 'departure_only':
                mark_label = 'Только уход (нет прихода)'
                time_str = _minutes_to_hhmm(d.departure_min)
            else:
                continue
            ws3.append([d.fio, _format_date(d.date), mark_label, time_str])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
